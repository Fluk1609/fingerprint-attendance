import io
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_excel_datetime(value):
    if not value:
        return "-"

    try:
        if isinstance(value, str):
            value = datetime.fromisoformat(
                value.replace("Z", "+00:00")
            )

        # ถ้าไม่มี timezone ให้ถือว่าเป็นเวลาไทยอยู่แล้ว
        if value.tzinfo is None:
            value = value.replace(
                tzinfo=timezone(timedelta(hours=7))
            )
        else:
            # แปลงเวลาเป็น UTC+7
            value = value.astimezone(
                timezone(timedelta(hours=7))
            )

        return value.strftime("%d/%m/%Y : %H:%M")

    except (ValueError, TypeError):
        return str(value)


HEADER_FILL = PatternFill(start_color="7031A0", end_color="7031A0", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14)
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)


def _style_header_row(ws, row_num, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _autofit_columns(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_excel_report(logs, start_date, end_date):
    """
    logs: list ของ sqlite3.Row ที่มีคอลัมน์
          scanned_at, status, fingerprint_id, student_id, full_name, department
    คืนค่า: bytes ของไฟล์ .xlsx
    """
    wb = Workbook()

    # ---------------- ชีต 1: สรุปรายบุคคล ----------------
    ws1 = wb.active
    ws1.title = "สรุปรายบุคคล"

    ws1.merge_cells("A1:F1")
    ws1["A1"] = f"รายงานสรุปการเช็คชื่อด้วยลายนิ้วมือ  ({start_date} ถึง {end_date})"
    ws1["A1"].font = TITLE_FONT
    ws1["A1"].alignment = Alignment(horizontal="left")

    headers = ["ลำดับ", "รหัสนักศึกษา/บุคลากร", "ชื่อ-นามสกุล", "หน่วยงาน/สาขา", "จำนวนครั้งที่สแกน", "สแกนล่าสุด"]
    header_row = 3
    for col, text in enumerate(headers, start=1):
        ws1.cell(row=header_row, column=col, value=text)
    _style_header_row(ws1, header_row, len(headers))

    # รวมข้อมูลตามคน
    summary = {}
    for row in logs:
        key = row["student_id"] or f"[ไม่ทราบผู้ใช้] {row['fingerprint_id']}"
        if key not in summary:
            summary[key] = {
                "full_name": row["full_name"] or "-",
                "department": row["department"] or "-",
                "count": 0,
                "last": row["scanned_at"],
            }
        summary[key]["count"] += 1
        if row["scanned_at"] > summary[key]["last"]:
            summary[key]["last"] = row["scanned_at"]

    r = header_row + 1
    for idx, (student_id, info) in enumerate(sorted(summary.items()), start=1):
        values = [
    idx,
    student_id,
    info["full_name"],
    info["department"],
    info["count"],
    format_excel_datetime(info["last"]),
]
        for col, val in enumerate(values, start=1):
            cell = ws1.cell(row=r, column=col, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col in (1, 5) else "left")
        r += 1

    if not summary:
        ws1.cell(row=r, column=1, value="ไม่มีข้อมูลในช่วงวันที่ที่เลือก").font = BODY_FONT

    _autofit_columns(ws1, [8, 22, 28, 22, 18, 22])

    # ---------------- ชีต 2: รายละเอียดทั้งหมด ----------------
    ws2 = wb.create_sheet("รายละเอียดทั้งหมด")
    headers2 = ["ลำดับ", "วันเวลาที่สแกน", "รหัสนักศึกษา/บุคลากร", "ชื่อ-นามสกุล", "หน่วยงาน/สาขา", "สถานะ"]
    for col, text in enumerate(headers2, start=1):
        ws2.cell(row=1, column=col, value=text)
    _style_header_row(ws2, 1, len(headers2))

    status_th = {"matched": "พบข้อมูล", "unmatched": "ไม่พบผู้ใช้ในระบบ"}

    for idx, row in enumerate(logs, start=1):
        values = [
            idx,
            format_excel_datetime(row["scanned_at"]),
            row["student_id"] or "-",
            row["full_name"] or "-",
            row["department"] or "-",
            status_th.get(row["status"], row["status"]),
        ]
        for col, val in enumerate(values, start=1):
            cell = ws2.cell(row=idx + 1, column=col, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col in (1, 6) else "left")

    if not logs:
        ws2.cell(row=2, column=1, value="ไม่มีข้อมูลในช่วงวันที่ที่เลือก").font = BODY_FONT

    _autofit_columns(ws2, [8, 20, 22, 28, 22, 20])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
