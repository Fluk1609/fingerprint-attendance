from openpyxl import load_workbook
from io import BytesIO
import threading
import queue
import base64
import time
import uuid
from pyzkfp import ZKFP2
from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_file,
    redirect,
    url_for,
    flash,
)

from asgiref.wsgi import WsgiToAsgi

import database as db
from report import build_excel_report
from datetime import datetime, date, timedelta, timezone, time as datetime_time

THAI_TZ = timezone(timedelta(hours=7))

app = Flask(__name__)
app.secret_key = "change-this-secret-key"

@app.template_filter("datetime_format_th")
def datetime_format_th(value):
    if not value:
        return "-"

    if isinstance(value, str):
        value = datetime.fromisoformat(
            value.replace("Z", "+00:00")
        )
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    value = value.astimezone(THAI_TZ)
    return value.strftime("%d/%m/%Y : %H:%M")


@app.template_filter("datetime_format")
def datetime_format(value):
    if not value:
        return "-"

    try:
        if isinstance(value, str):
            value = datetime.fromisoformat(
                value.replace("Z", "+00:00")
            )

        return value.strftime("%d/%m/%Y : %H:%M")

    except (ValueError, TypeError):
        return str(value)

def delete_fingerprint_template(fingerprint_id):
    """ล้าง Template ใน Supabase (ไม่ใช้ fingerprint_templates.json)"""
    try:
        db.supabase.table("users").update(
            {"fingerprint_template": None}
        ).eq("fingerprint_id", str(fingerprint_id)).execute()
        print(f"[SCANNER] ล้าง Template รหัส {fingerprint_id} จาก Supabase แล้ว")
    except Exception as exc:
        print(f"[SCANNER] ล้าง Template ใน Supabase ไม่สำเร็จ: {exc}")
        raise

db.init_db()

scanner_queue = queue.Queue()
scanner_status = {
    "ready": False,
    "error": None
}

scanner_lock = threading.Lock()

enrollment_jobs = {}
enrollment_lock = threading.Lock()


active_subject_id = None
active_subject_lock = threading.Lock()

def get_current_subject(now=None):
    now = now or datetime.now(THAI_TZ)

    if now.tzinfo is None:
        now = now.replace(tzinfo=THAI_TZ)
    else:
        now = now.astimezone(THAI_TZ)

    response = (
        db.supabase.table("subjects")
        .select(
            "id, code, name, class_days, "
            "start_time, end_time, late_after_minutes"
        )
        .execute()
    )

    weekday = now.strftime("%A")
    current_time = now.time().replace(tzinfo=None)

    matches = []

    for subject in response.data or []:
        class_days = subject.get("class_days") or []

        if isinstance(class_days, str):
            class_days = [
                day.strip()
                for day in class_days.split(",")
                if day.strip()
            ]

        class_days = [
            str(day).strip().capitalize()
            for day in class_days
        ]

        print(
            f"[SUBJECT DEBUG] วิชา={subject.get('name')}, "
            f"วันเรียน={class_days}, วันปัจจุบัน={weekday}"
        )

        if weekday not in class_days:
            continue

        try:
            start_text = str(subject.get("start_time") or "").strip()
            end_text = str(subject.get("end_time") or "").strip()

            start_time = datetime.strptime(
                start_text[:8],
                "%H:%M:%S" if len(start_text) >= 8 else "%H:%M"
            ).time()

            end_time = datetime.strptime(
                end_text[:8],
                "%H:%M:%S" if len(end_text) >= 8 else "%H:%M"
            ).time()

        except ValueError as error:
            print(
                f"[SUBJECT DEBUG] เวลาไม่ถูกต้อง "
                f"วิชา={subject.get('name')}: {error}"
            )
            continue

        print(
            f"[SUBJECT DEBUG] วิชา={subject.get('name')}, "
            f"เวลาเรียน={start_time}-{end_time}, "
            f"เวลาปัจจุบัน={current_time}"
        )

        if start_time <= current_time <= end_time:
            matches.append(subject)

    if len(matches) == 1:
        return matches[0]

    if len(matches) > 1:
        print(
            "[SUBJECT DEBUG] พบวิชาที่เวลาเรียนทับซ้อนกัน: "
            + ", ".join(str(x.get("name")) for x in matches)
        )
        matches.sort(
            key=lambda subject: str(subject.get("start_time") or ""),
            reverse=True
        )

        selected = matches[0]

        print(
            f"[SUBJECT DEBUG] เลือกวิชา: {selected.get('name')} "
            f"เวลาเริ่ม {selected.get('start_time')}"
        )

        return selected

    return None

def current_subject_id():
    subject = get_current_subject()
    return subject.get("id") if subject else None


@app.route("/")
def dashboard():
    today = datetime.now(THAI_TZ).date()
    start_of_day = datetime.combine(
        today, datetime.min.time(), tzinfo=THAI_TZ
    )
    start_of_tomorrow = datetime.combine(
        today + timedelta(days=1),
        datetime.min.time(),
        tzinfo=THAI_TZ
    )
    start_utc = start_of_day.astimezone(timezone.utc)
    end_utc = start_of_tomorrow.astimezone(timezone.utc)
    response = (
        db.supabase.table("scan_logs")
        .select("""
            id, scanned_at, status, fingerprint_id, subject_id,
            users (student_id, full_name, department),
            subjects (id, code, name)
        """)
        .gte("scanned_at", start_utc.isoformat())
        .lt("scanned_at", end_utc.isoformat())
        .order("scanned_at", desc=True)
        .limit(50)
        .execute()
    )

    logs = []
    for row in response.data or []:
        user = row.get("users") or {}
        subject = row.get("subjects") or {}

        logs.append({
            "id": row.get("id"),
            "scanned_at": row.get("scanned_at"),
            "status": row.get("status"),
            "fingerprint_id": row.get("fingerprint_id"),
            "student_id": user.get("student_id"),
            "full_name": user.get("full_name"),
            "department": user.get("department"),
            "subject_id": row.get("subject_id"),
            "subject_code": subject.get("code"),
            "subject_name": subject.get("name"),
        })

    active_subject = get_current_subject()

    return render_template(
        "index.html",
        logs=logs,
        active_subject=active_subject,
    )


@app.route("/api/current-subject")
def api_current_subject():
    subject = get_current_subject()
    if not subject:
        return jsonify({"active": False})
    return jsonify({
        "active": True,
        "id": subject.get("id"),
        "code": subject.get("code"),
        "name": subject.get("name"),
        "class_days": subject.get("class_days") or [],
        "start_time": str(subject.get("start_time") or "")[:5],
        "end_time": str(subject.get("end_time") or "")[:5],
        "late_after_minutes": subject.get("late_after_minutes") or 0,
    })


@app.route("/attendance/subject", methods=["POST"])
def set_active_subject():
    global active_subject_id
    subject_id_raw = (request.form.get("subject_id") or "").strip()

    if not subject_id_raw:
        with active_subject_lock:
            active_subject_id = None
        print("[SUBJECT DEBUG] ยกเลิกรายวิชาแล้ว: active_subject_id=None")
        flash("ยกเลิกรายวิชาที่กำลังเช็กชื่อแล้ว", "success")
        return redirect(url_for("dashboard"))

    try:
        subject_id = int(subject_id_raw)
        response = (
            db.supabase.table("subjects")
            .select("id, name")
            .eq("id", subject_id)
            .limit(1)
            .execute()
        )
        subject = (response.data or [None])[0]

        if not subject:
            flash("ไม่พบรายวิชาที่เลือก", "error")
        else:
            with active_subject_lock:
                active_subject_id = subject_id
            print(
                f"[SUBJECT DEBUG] เริ่มเช็กชื่อ: "
                f"active_subject_id={active_subject_id}, "
                f"name={subject.get('name')}"
            )
            flash(f"เริ่มเช็กชื่อรายวิชา {subject.get('name')}", "success")
    except (TypeError, ValueError):
        flash("รหัสรายวิชาไม่ถูกต้อง", "error")
    except Exception as exc:
        flash(f"กำหนดรายวิชาที่เช็กชื่อไม่สำเร็จ: {exc}", "error")

    return redirect(url_for("dashboard"))


@app.route("/api/scans/recent")
def api_recent_scans():
    logs = db.get_recent_logs(limit=50)

    return jsonify([
        dict(row)
        for row in logs
    ])


@app.route("/api/scan", methods=["POST"])
def api_scan():
    payload = request.get_json(silent=True) or {}

    fingerprint_id = str(
        payload.get("fingerprint_id") or ""
    ).strip()

    scanned_at = payload.get("scanned_at") or datetime.now(
        THAI_TZ
    ).isoformat(timespec="seconds")

    if not fingerprint_id:
        return jsonify({
            "error": "ต้องระบุ fingerprint_id"
        }), 400

    user = db.get_user_by_fingerprint(fingerprint_id)

    scan_dt = datetime.fromisoformat(str(scanned_at).replace("Z", "+00:00"))
    if scan_dt.tzinfo is None:
        scan_dt = scan_dt.replace(tzinfo=THAI_TZ)
    subject = get_current_subject(scan_dt)
    subject_id = subject.get("id") if subject else None

    print(
        f"[SCAN DEBUG] fingerprint_id={fingerprint_id}, "
        f"subject_id={subject_id}, "
        f"scanned_at={scanned_at}"
    )

    if user:
        attendance_status = None

        if subject_id is not None:
            attendance_status = get_attendance_status(
                subject_id,
                scanned_at
            )

        print(
            f"[ATTENDANCE DEBUG] "
            f"student_id={user.get('student_id')}, "
            f"attendance_status={attendance_status}"
        )

        log_status = attendance_status or "no_class"

        db.add_scan_log(
            fingerprint_id,
            user["id"],
            log_status,
            scanned_at,
            subject_id=subject_id,
        )

        return jsonify({
            "status": log_status,
            "attendance_status": attendance_status,
            "student_id": user.get("student_id"),
            "full_name": user.get("full_name"),
            "department": user.get("department"),
            "subject_id": subject_id,
            "scanned_at": scanned_at,
        }), 200

    db.add_scan_log(
        fingerprint_id,
        None,
        "unmatched",
        scanned_at,
        subject_id=subject_id,
    )

    return jsonify({
        "status": "unmatched",
        "attendance_status": None,
        "subject_id": subject_id,
        "fingerprint_id": fingerprint_id,
        "scanned_at": scanned_at,
    }), 200

@app.route("/users")
def users_page():
    users = db.get_all_users()
    sections = _master_list("sections", "name")

    for user in users:
        section_id = user.get("section_id")
        section = next(
            (item for item in sections if str(item.get("id")) == str(section_id)),
            None,
        )
        user["section_name"] = section.get("name") if section else "-"

    return render_template(
        "users.html",
        users=users,
        sections=sections,
    )


@app.route("/users/add", methods=["POST"])
def users_add():

    student_id = request.form.get(
        "student_id",
        ""
    ).strip()

    full_name = request.form.get(
        "full_name",
        ""
    ).strip()

    department = request.form.get(
        "department",
        ""
    ).strip()

    fingerprint_id = request.form.get(
        "fingerprint_id",
        ""
    ).strip()

    section_id_raw = (request.form.get("section_id") or "").strip()
    section_name = (request.form.get("section") or "").strip()
    section_id = None

    if section_id_raw:
        try:
            section_id = int(section_id_raw)
        except ValueError:
            flash("Section ไม่ถูกต้อง", "error")
            return redirect(url_for("users_page"))
    elif section_name:
        response = db.supabase.table("sections").select("id,name").execute()
        matched = next(
            (item for item in (response.data or [])
             if item.get("name", "").strip().casefold() == section_name.casefold()),
            None,
        )
        if not matched:
            flash(f"ไม่พบ Section '{section_name}' ใน Master Section", "error")
            return redirect(url_for("users_page"))
        section_id = int(matched["id"])

    if not (
        student_id
        and full_name
        and fingerprint_id
    ):

        flash(
            "กรุณากรอกรหัส, ชื่อ-นามสกุล "
            "และรหัสลายนิ้วมือให้ครบ",
            "error",
        )

        return redirect(
            url_for("users_page")
        )

    try:

        created_at = datetime.now(
            timezone(timedelta(hours=7))
        ).isoformat(timespec="seconds")

        db.add_user(
            student_id,
            full_name,
            department,
            fingerprint_id,
            created_at=created_at,
            section_id=section_id,
        )

        flash(
            f"เพิ่มผู้ใช้ {full_name} สำเร็จ",
            "success",
        )

    except Exception as exc:
        print(f"[ERROR] เพิ่มผู้ใช้ไม่สำเร็จ: {exc}")

        flash(
            f"เพิ่มผู้ใช้ไม่สำเร็จ: "
            f"รหัสนักศึกษาหรือรหัสลายนิ้วมือซ้ำ ({exc})",
            "error",
        )
    return redirect(
        url_for("users_page")
    )



@app.route("/users/import-excel", methods=["POST"])
def users_import_excel():
    file = request.files.get("excel_file")

    if not file or not file.filename:
        flash("กรุณาเลือกไฟล์ Excel", "error")
        return redirect(url_for("users_page"))

    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("รองรับไฟล์ .xlsx หรือ .xlsm เท่านั้น", "error")
        return redirect(url_for("users_page"))

    workbook = None

    try:
        workbook = load_workbook(
            file,
            read_only=True,
            data_only=True
        )
        sheet = workbook.active

        header_values = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            None
        )

        if not header_values:
            raise ValueError("ไฟล์ Excel ไม่มีหัวคอลัมน์")

        headers = [
            str(value or "").strip().lower()
            for value in header_values
        ]

        required = [
            "student_id",
            "full_name",
            "department",
            "section"
        ]

        missing = [name for name in required if name not in headers]

        if missing:
            raise ValueError(
                "ไม่พบคอลัมน์: " + ", ".join(missing)
            )

        section_response = (
            db.supabase.table("sections")
            .select("id,name")
            .execute()
        )

        section_map = {
            str(item["name"]).strip().casefold(): int(item["id"])
            for item in (section_response.data or [])
            if item.get("name") and item.get("id") is not None
        }

        rows = []
        row_errors = []
        seen_student_ids = set()

        for excel_row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2
        ):
            row = dict(zip(headers, values))

            if not any(value is not None for value in values):
                continue

            raw_student_id = row.get("student_id")
            full_name = str(row.get("full_name") or "").strip()
            department = str(row.get("department") or "").strip()
            section_name = str(row.get("section") or "").strip()

            student_id = str(raw_student_id or "").strip()
            if student_id.endswith(".0"):
                student_id = student_id[:-2]

            if not student_id or not full_name or not department:
                row_errors.append({
                    "row": excel_row_number,
                    "message": "ข้อมูล student_id, full_name หรือ department ไม่ครบ"
                })
                continue

            if not student_id.isdigit():
                row_errors.append({
                    "row": excel_row_number,
                    "message": f"student_id '{student_id}' ต้องเป็นตัวเลข"
                })
                continue

            if student_id in seen_student_ids:
                row_errors.append({
                    "row": excel_row_number,
                    "message": f"student_id '{student_id}' ซ้ำในไฟล์"
                })
                continue

            seen_student_ids.add(student_id)

            normalized_section = section_name.casefold()
            if normalized_section.startswith("section "):
                normalized_section = normalized_section[len("section "):].strip()

            section_id = None

            if normalized_section:
                section_id = section_map.get(section_name.casefold())

                if section_id is None and normalized_section.isdigit():
                    candidate_id = int(normalized_section)
                    if candidate_id in section_map.values():
                        section_id = candidate_id

                if section_id is None:
                    row_errors.append({
                        "row": excel_row_number,
                        "message": f"ไม่พบ Section '{section_name}' ใน Master Section"
                    })
                    continue

            else:
                row_errors.append({
                    "row": excel_row_number,
                    "message": "กรุณาระบุ section"
                })
                continue

            rows.append({
                "student_id": int(student_id),
                "full_name": full_name,
                "department": department,
                "section_id": section_id,
            })

        if not rows:
            flash(
                "ไม่มีข้อมูลที่สามารถนำเข้าได้"
                + (
                    f" พบข้อผิดพลาด {len(row_errors)} รายการ"
                    if row_errors else ""
                ),
                "error"
            )
            for error in row_errors[:10]:
                flash(
                    f"แถว {error['row']}: {error['message']}",
                    "error"
                )
            return redirect(url_for("users_page"))

        result = db.import_users_from_excel(rows)

        result.setdefault("inserted", 0)
        result.setdefault("updated", 0)
        result.setdefault("skipped", 0)
        result.setdefault("errors", [])

        all_errors = row_errors + result["errors"]

        flash(
            f"เพิ่มใหม่ {result['inserted']} คน, "
            f"อัปเดตข้อมูลเดิม {result['updated']} คน, "
            f"ข้าม {result['skipped']} คน, "
            f"ผิดพลาด {len(all_errors)} รายการ",
            "success" if not all_errors else "error"
        )

        for error in all_errors[:10]:
            flash(
                f"แถว {error.get('row', '?')}: "
                f"{error.get('message', 'เกิดข้อผิดพลาด')}",
                "error"
            )

    except Exception as exc:
        app.logger.exception("Import Excel failed")
        flash(f"นำเข้า Excel ไม่สำเร็จ: {exc}", "error")

    finally:
        if workbook is not None:
            workbook.close()

    return redirect(url_for("users_page"))



@app.route("/users/delete/<int:user_id>", methods=["POST"])
def users_delete(user_id):
    try:
        users = db.get_all_users()
        user = next(
            (u for u in users if int(u["id"]) == user_id),
            None,
        )
        if not user:
            flash("ไม่พบผู้ใช้ที่ต้องการลบ", "error")
            return redirect(url_for("users_page"))

        fingerprint_id = user.get("fingerprint_id")

        db.delete_user(user_id)

        if fingerprint_id:
            delete_fingerprint_template(fingerprint_id)

        flash("ลบผู้ใช้และ Template เรียบร้อยแล้ว", "success")

    except Exception as exc:
        print(f"[ERROR] ลบผู้ใช้ไม่สำเร็จ: {exc}")
        flash(f"ลบผู้ใช้ไม่สำเร็จ: {exc}", "error")

    return redirect(url_for("users_page"))

def _master_list(table, order_column):
    response = db.supabase.table(table).select("*").order(order_column).execute()
    return response.data or []

@app.route("/masters/sections", methods=["GET", "POST"])
def sections_page():
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("กรุณากรอกชื่อ Section", "error")
        else:
            try:
                db.supabase.table("sections").insert({"name": name}).execute()
                flash("เพิ่ม Section สำเร็จ", "success")
            except Exception as exc:
                flash(f"เพิ่ม Section ไม่สำเร็จ: {exc}", "error")
        return redirect(url_for("sections_page"))
    try:
        sections = _master_list("sections", "name")
    except Exception as exc:
        sections = []
        flash("อ่านตาราง sections ไม่สำเร็จ กรุณารัน SQL สำหรับ Master Data ก่อน: " + str(exc), "error")
    return render_template("sections.html", sections=sections)

@app.route("/masters/sections/<int:section_id>/delete", methods=["POST"])
def sections_delete(section_id):
    try:
        db.supabase.table("sections").delete().eq("id", section_id).execute()
        flash("ลบ Section สำเร็จ", "success")
    except Exception as exc:
        flash(f"ลบ Section ไม่สำเร็จ: {exc}", "error")
    return redirect(url_for("sections_page"))

@app.route("/masters/subjects", methods=["GET", "POST"])
def subjects_page():
    if request.method == "POST":
        try:
            code = request.form.get("code", "").strip() or None
            name = request.form.get("name", "").strip()
            section_id = request.form.get("section_id")
            class_days = request.form.getlist("class_days")
            start_time = request.form.get("start_time")
            end_time = request.form.get("end_time")
            late_after_minutes = request.form.get("late_after_minutes", "0")

            if not name or not section_id:
                flash("กรุณากรอกชื่อวิชาและเลือก Section", "error")
                return redirect(url_for("subjects_page"))

            payload = {
                "code": code,
                "name": name,
                "section_id": int(section_id),
                "class_days": class_days,
                "start_time": start_time,
                "end_time": end_time,
                "late_after_minutes": int(late_after_minutes),
            }

            result = (
                db.supabase
                .table("subjects")
                .insert(payload)
                .execute()
            )

            print("เพิ่มวิชาสำเร็จ:", result.data)
            flash("เพิ่มวิชาเรียบร้อยแล้ว", "success")

        except Exception as e:
            print("เพิ่มวิชาไม่สำเร็จ:", repr(e))
            flash(f"เพิ่มวิชาไม่สำเร็จ: {e}", "error")

        return redirect(url_for("subjects_page"))

    subjects_result = (
        db.supabase
        .table("subjects")
        .select("*, sections(name)")
        .order("id")
        .execute()
    )

    sections_result = (
        db.supabase
        .table("sections")
        .select("*")
        .order("id")
        .execute()
    )

    return render_template(
        "subjects.html",
        subjects=subjects_result.data or [],
        sections=sections_result.data or [],
    )


@app.route("/masters/subjects/edit/<int:subject_id>", methods=["GET", "POST"])
def subjects_edit(subject_id):
    try:
        if request.method == "POST":
            code = request.form.get("code", "").strip()
            name = request.form.get("name", "").strip()
            section_id = request.form.get("section_id", "").strip()
            class_days = request.form.getlist("class_days")
            start_time = request.form.get("start_time", "").strip()
            end_time = request.form.get("end_time", "").strip()
            late_after_minutes = request.form.get("late_after_minutes", "0").strip()

            if not code or not name or not section_id or not start_time or not end_time:
                flash("กรุณากรอกข้อมูลให้ครบถ้วน", "error")
                return redirect(url_for("subjects_edit", subject_id=subject_id))

            try:
                section_id = int(section_id)
                late_after_minutes = int(late_after_minutes)

                if late_after_minutes < 0:
                    raise ValueError("นาทีผ่อนผันต้องไม่ติดลบ")

                for time_value in (start_time, end_time):
                    hour, minute = map(int, time_value.split(":"))
                    if not (0 <= hour <= 23 and 0 <= minute <= 59):
                        raise ValueError("รูปแบบเวลาไม่ถูกต้อง")

            except (ValueError, TypeError):
                flash("กรุณาตรวจสอบ Section เวลาเรียน และนาทีผ่อนผัน", "error")
                return redirect(url_for("subjects_edit", subject_id=subject_id))

            payload = {
                "code": code,
                "name": name,
                "section_id": section_id,
                "class_days": class_days,
                "start_time": start_time,
                "end_time": end_time,
                "late_after_minutes": late_after_minutes,
            }

            result = (
                db.supabase
                .table("subjects")
                .update(payload)
                .eq("id", subject_id)
                .execute()
            )

            if not result.data:
                flash("ไม่พบรายวิชาที่ต้องการแก้ไข", "error")
                return redirect(url_for("subjects_page"))

            flash("แก้ไขรายวิชาเรียบร้อยแล้ว", "success")
            return redirect(url_for("subjects_page"))

        subject_result = (
            db.supabase
            .table("subjects")
            .select("*")
            .eq("id", subject_id)
            .execute()
        )

        if not subject_result.data:
            flash("ไม่พบรายวิชาที่ต้องการแก้ไข", "error")
            return redirect(url_for("subjects_page"))

        sections_result = (
            db.supabase
            .table("sections")
            .select("*")
            .order("id")
            .execute()
        )

        return render_template(
            "subjects_edit.html",
            subject=subject_result.data[0],
            sections=sections_result.data or [],
        )

    except Exception as e:
        print(f"แก้ไขวิชาไม่สำเร็จ: {repr(e)}")
        flash(f"แก้ไขวิชาไม่สำเร็จ: {e}", "error")
        return redirect(url_for("subjects_page"))


@app.route("/masters/subjects/delete/<int:subject_id>", methods=["POST"])
def subjects_delete(subject_id):
    try:
        db.supabase.table("scan_logs") \
            .delete() \
            .eq("subject_id", subject_id) \
            .execute()

        db.supabase.table("subjects") \
            .delete() \
            .eq("id", subject_id) \
            .execute()

        flash("ลบวิชาและประวัติการสแกนเรียบร้อยแล้ว", "success")

    except Exception as e:
        print(f"ลบวิชาไม่สำเร็จ: {e}")
        flash(f"ลบวิชาไม่สำเร็จ: {e}", "error")

    return redirect(url_for("subjects_page"))


@app.route("/masters/subjects/delete/<int:subject_id>", methods=["POST"])
def delete_subject(subject_id):
    try:
        db.supabase.table("scan_logs") \
            .delete() \
            .eq("subject_id", subject_id) \
            .execute()
        db.supabase.table("subjects") \
            .delete() \
            .eq("id", subject_id) \
            .execute()
        flash("ลบวิชาและประวัติการสแกนเรียบร้อยแล้ว", "success")
    except Exception as e:
        print(f"ลบวิชาไม่สำเร็จ: {e}")
        flash("ลบวิชาไม่สำเร็จ", "error")

    return redirect(url_for("subjects_page"))

@app.route("/reports")
def reports_page():

    today = date.today()

    start_default = (
        today - timedelta(days=7)
    ).isoformat()

    end_default = today.isoformat()

    start_date = request.args.get(
        "start_date",
        start_default,
    )

    end_date = request.args.get(
        "end_date",
        end_default,
    )

    user_id = (
        request.args.get("user_id")
        or None
    )

    logs = db.get_logs_in_range(
        start_date,
        end_date,
        user_id,
    )

    users = db.get_all_users()

    summary = {}

    for row in logs:

        key = (
            row["student_id"]
            or "ไม่พบข้อมูลผู้ใช้"
        )

        if key not in summary:

            summary[key] = {
                "full_name": (
                    row["full_name"]
                    or "-"
                ),

                "department": (
                    row["department"]
                    or "-"
                ),

                "count": 0,

                "first": row["scanned_at"],

                "last": row["scanned_at"],
            }

        summary[key]["count"] += 1

        summary[key]["last"] = (
            row["scanned_at"]
        )

    return render_template(
        "reports.html",
        logs=logs,
        users=users,
        summary=summary,
        start_date=start_date,
        end_date=end_date,
        selected_user_id=user_id,
    )


@app.route("/reports/export")
def reports_export():

    today = date.today()

    start_date = request.args.get(
        "start_date",
        (
            today - timedelta(days=7)
        ).isoformat(),
    )

    end_date = request.args.get(
        "end_date",
        today.isoformat(),
    )

    user_id = (
        request.args.get("user_id")
        or None
    )

    logs = db.get_logs_in_range(
        start_date,
        end_date,
        user_id,
    )

    excel_bytes = build_excel_report(
        logs,
        start_date,
        end_date,
    )

    filename = (
        f"attendance_report_"
        f"{start_date}_to_{end_date}.xlsx"
    )

    return send_file(
        BytesIO(excel_bytes),
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=filename,
    )

@app.route("/api/enroll", methods=["POST"])
def api_enroll():
    payload = request.get_json(silent=True) or {}
    fingerprint_id = str(payload.get("fingerprint_id", "")).strip()
    replace_existing = payload.get("replace_existing") is True

    if not fingerprint_id.isdigit():
        return jsonify({
            "error": "รหัสลายนิ้วมือต้องเป็นตัวเลข"
        }), 400

    try:
        has_template = db.has_fingerprint_template(fingerprint_id)
    except AttributeError:
        return jsonify({
            "error": "database.py ยังไม่มีฟังก์ชัน has_fingerprint_template()"
        }), 500
    except Exception as exc:
        print(f"[ENROLL] ตรวจสอบ Template ไม่สำเร็จ: {exc}")
        return jsonify({
            "error": "ตรวจสอบลายนิ้วมือเดิมใน Supabase ไม่สำเร็จ"
        }), 500

    if has_template and not replace_existing:
        return jsonify({
            "requires_confirmation": True,
            "fingerprint_id": fingerprint_id,
            "message": "ผู้ใช้นี้มีลายนิ้วมืออยู่แล้ว ต้องการเปลี่ยนลายนิ้วมือหรือไม่?"
        }), 409

    with enrollment_lock:
        if any(
            job["status"] in ("pending", "processing")
            for job in enrollment_jobs.values()
        ):
            return jsonify({
                "error": "มีงานลงทะเบียนที่กำลังดำเนินการอยู่"
            }), 409

        job_id = str(uuid.uuid4())

        job = {
            "job_id": job_id,
            "fingerprint_id": fingerprint_id,
            "replace_existing": replace_existing,
            "status": "pending",
            "message": "รอเครื่องสแกนรับงาน"
        }

        enrollment_jobs[job_id] = job
        scanner_queue.put(job.copy())

    return jsonify({
        "job_id": job_id,
        "status": "pending",
        "message": "ส่งคำขอลงทะเบียนแล้ว"
    }), 202



@app.route("/api/enroll/<job_id>", methods=["GET"])
def api_enroll_status(job_id):
    with enrollment_lock:
        job = enrollment_jobs.get(job_id)

        if not job:
            return jsonify({"error": "ไม่พบงานลงทะเบียน"}), 404

        return jsonify({
            "job_id": job["job_id"],
            "fingerprint_id": job["fingerprint_id"],
            "status": job["status"],
            "message": job["message"]
        })


@app.route("/api/scanner/enrollment/next", methods=["GET"])
def api_scanner_next_enrollment():
    with enrollment_lock:
        for job in enrollment_jobs.values():
            if job["status"] == "pending":
                job["status"] = "processing"
                job["message"] = "กรุณาวางนิ้วบนเครื่องสแกน 3 ครั้ง"
                return jsonify(job)

    return jsonify({"job": None})


@app.route("/api/scanner/enrollment/result", methods=["POST"])
def api_scanner_enrollment_result():
    payload = request.get_json(silent=True) or {}
    job_id = payload.get("job_id")
    success = payload.get("success") is True
    message = payload.get("message", "")

    with enrollment_lock:
        job = enrollment_jobs.get(job_id)

        if not job or job["status"] != "processing":
            return jsonify({"error": "ไม่พบงานหรือสถานะงานไม่ถูกต้อง"}), 404

        job["status"] = "success" if success else "failed"
        job["message"] = (
            message or (
                "ลงทะเบียนลายนิ้วมือสำเร็จ"
                if success else "ลงทะเบียนไม่สำเร็จ"
            )
        )

    return jsonify({"status": "ok"})

@app.route("/api/enroll/reset", methods=["POST"])
def api_enroll_reset():
    with enrollment_lock:
        for job in enrollment_jobs.values():
            if job["status"] in ("pending", "processing"):
                job["status"] = "failed"
                job["message"] = "ยกเลิกงานลงทะเบียนเดิมแล้ว"

    return jsonify({
        "status": "success",
        "message": "ยกเลิกงานลงทะเบียนที่ค้างแล้ว"
    })


def get_attendance_status(subject_id, scanned_at):
    if subject_id is None or str(subject_id).strip() == "":
        return None

    response = (
        db.supabase.table("subjects")
        .select("id, name, start_time, late_after_minutes")
        .eq("id", int(subject_id))
        .limit(1)
        .execute()
    )
    subjects = response.data or []
    if not subjects:
        print(f"[ATTENDANCE DEBUG] ไม่พบ subject_id={subject_id}")
        return None

    subject = subjects[0]
    start_time_value = subject.get("start_time")
    if not start_time_value:
        print(
            f"[ATTENDANCE DEBUG] วิชา {subject_id} "
            "ไม่มี start_time"
        )
        return None
    start_time_text = str(start_time_value).strip()
    parsed_start_time = None
    for time_format in ("%H:%M:%S", "%H:%M"):
        try:
            parsed_start_time = datetime.strptime(
                start_time_text[:8] if time_format == "%H:%M:%S"
                else start_time_text[:5],
                time_format
            ).time()
            break
        except ValueError:
            continue

    if parsed_start_time is None:
        print(
            f"[ATTENDANCE ERROR] รูปแบบ start_time ไม่ถูกต้อง: "
            f"{start_time_value!r}"
        )
        return None

    try:
        scan_datetime = datetime.fromisoformat(
            str(scanned_at).replace("Z", "+00:00")
        )
    except (TypeError, ValueError):
        print(f"[ATTENDANCE ERROR] scanned_at ไม่ถูกต้อง: {scanned_at!r}")
        return None

    if scan_datetime.tzinfo is None:
        scan_datetime = scan_datetime.replace(tzinfo=THAI_TZ)
    else:
        scan_datetime = scan_datetime.astimezone(THAI_TZ)

    start_datetime = datetime.combine(
        scan_datetime.date(),
        parsed_start_time,
        tzinfo=THAI_TZ
    )

    try:
        grace_minutes = max(0, int(subject.get("late_after_minutes") or 0))
    except (TypeError, ValueError):
        grace_minutes = 0

    deadline = start_datetime + timedelta(minutes=grace_minutes)
    result = "on_time" if scan_datetime <= deadline else "late"

    print(
        f"[ATTENDANCE DEBUG] subject_id={subject_id}, "
        f"subject={subject.get('name')}, "
        f"start={start_datetime.isoformat()}, "
        f"deadline={deadline.isoformat()}, "
        f"scan={scan_datetime.isoformat()}, "
        f"result={result}"
    )
    return result


def scanner_worker():
    zk = ZKFP2()
    initialized = False
    opened = False
    templates = {}
    last_id = None
    last_time = 0.0

    def capture_template():
        while True:
            result = zk.AcquireFingerprint()

            if isinstance(result, (tuple, list)) and len(result) >= 2:
                template, _image = result

                if template:
                    return bytes(template)

            time.sleep(0.2)

    def enroll_with_sdk(job):
        job_id = job["job_id"]
        fingerprint_id = str(job["fingerprint_id"])
        fid = int(fingerprint_id)

        with enrollment_lock:
            enrollment_jobs[job_id]["status"] = "processing"
            enrollment_jobs[job_id]["message"] = (
                "กรุณาวางนิ้วบนเครื่องสแกน"
            )

        try:
            existing_template = templates.get(fingerprint_id)
            replace_existing = job.get("replace_existing") is True

            if existing_template and not replace_existing:
                raise ValueError(
                    "รหัสนี้มี Template อยู่แล้ว ต้องยืนยันเปลี่ยนลายนิ้วมือก่อน"
                )

            captures = []

            for index in range(3):
                with enrollment_lock:
                    enrollment_jobs[job_id]["message"] = (
                        f"กรุณาวางนิ้วครั้งที่ {index + 1}/3"
                    )

                print(
                    f"[SCANNER] รหัส {fingerprint_id}: "
                    f"วางนิ้วครั้งที่ {index + 1}/3"
                )

                captures.append(capture_template())
                print("[SCANNER] อ่านสำเร็จ กรุณานำนิ้วออก")
                time.sleep(1)

            merged, merged_length = zk.DBMerge(
                captures[0],
                captures[1],
                captures[2]
            )

            if not merged:
                raise RuntimeError("รวม Template ไม่สำเร็จ")

            merged = bytes(merged)

            if merged_length and len(merged) != int(merged_length):
                raise RuntimeError("ความยาว Template ไม่ถูกต้อง")
            template_base64 = base64.b64encode(merged).decode("ascii")
            db.update_fingerprint_template(fingerprint_id, template_base64)

            try:
                try:
                    zk.DBDel(fid)
                except Exception:
                    pass

                zk.DBAdd(fid, merged)
                templates[fingerprint_id] = merged
            except Exception:
                try:
                    zk.DBDel(fid)
                except Exception:
                    pass

                if existing_template:
                    try:
                        zk.DBAdd(fid, existing_template)
                        templates[fingerprint_id] = existing_template
                    except Exception as restore_exc:
                        print(
                            f"[SCANNER] คืน Template เดิมใน SDK ไม่สำเร็จ: "
                            f"{restore_exc}"
                        )
                else:
                    templates.pop(fingerprint_id, None)
                raise

            with enrollment_lock:
                enrollment_jobs[job_id]["status"] = "success"
                enrollment_jobs[job_id]["message"] = (
                    "ลงทะเบียนลายนิ้วมือสำเร็จ"
                )

            print(f"[SCANNER] ลงทะเบียน {fingerprint_id} สำเร็จ")

        except Exception as exc:
            with enrollment_lock:
                enrollment_jobs[job_id]["status"] = "failed"
                enrollment_jobs[job_id]["message"] = str(exc)

            print(f"[SCANNER] ลงทะเบียนไม่สำเร็จ: {exc}")

    try:
        zk.Init()
        initialized = True

        if not zk.GetDeviceCount():
            raise RuntimeError("ไม่พบเครื่องสแกน ZK9500")

        zk.OpenDevice(0)
        opened = True

        loaded_templates = db.get_all_fingerprint_templates() or {}
        templates = {}

        for fingerprint_id, template in loaded_templates.items():
            if template is None:
                continue

            if isinstance(template, str):
                template_bytes = base64.b64decode(template, validate=True)
            else:
                template_bytes = bytes(template)

            normalized_id = str(fingerprint_id).strip()

            if not normalized_id.isdigit():
                raise ValueError(
                    f"Fingerprint ID ต้องเป็นตัวเลข: {normalized_id}"
                )

            if not template_bytes:
                raise ValueError(
                    f"Template ว่าง: fingerprint_id={normalized_id}"
                )

            try:
                zk.DBAdd(int(normalized_id), template_bytes)
                templates[normalized_id] = template_bytes
            except Exception as exc:
                print(
                    f"[SCANNER ERROR] โหลด Template ไม่สำเร็จ: "
                    f"id={normalized_id}, "
                    f"bytes={len(template_bytes)}, "
                    f"error={exc!r}"
                )
                raise

        scanner_status["ready"] = True
        scanner_status["error"] = None

        while True:
            try:
                job = scanner_queue.get_nowait()
            except queue.Empty:
                job = None

            if job:
                try:
                    enroll_with_sdk(job)
                finally:
                    scanner_queue.task_done()

                last_id = None
                last_time = 0.0
                continue

            result = zk.AcquireFingerprint()

            if not isinstance(result, (tuple, list)) or len(result) < 2:
                time.sleep(0.2)
                continue

            template, _image = result

            if not template:
                time.sleep(0.2)
                continue

            try:
                uid, score = zk.DBIdentify(bytes(template))

                if uid is None or int(uid) < 0:
                    continue

                fingerprint_id = str(uid)
                now = time.time()

                if (
                    fingerprint_id != last_id
                    or now - last_time >= 3
                ):
                    user = db.get_user_by_fingerprint(fingerprint_id)

                    scanned_at = datetime.now(THAI_TZ).isoformat(timespec="seconds")
                    active_subject = get_current_subject(datetime.now(THAI_TZ))
                    subject_id = active_subject.get("id") if active_subject else None

                    if user:
                        if subject_id is None:
                            attendance_status = None
                            log_status = "no_class"
                        else:
                            attendance_status = get_attendance_status(subject_id, scanned_at)
                            log_status = attendance_status or "no_class"
                    else:
                        attendance_status = None
                        log_status = "unmatched"

                    print(
                        f"[SCAN DEBUG] fingerprint_id={fingerprint_id}, "
                        f"subject_id={subject_id}, "
                        f"attendance_status={attendance_status}, "
                        f"log_status={log_status}"
                    )

                    db.add_scan_log(
                        fingerprint_id,
                        user["id"] if user else None,
                        log_status,
                        scanned_at,
                        subject_id=subject_id,
                    )

                    print(
                        f"[SCANNER] {fingerprint_id}: "
                        f"{user['full_name'] if user else 'ไม่พบผู้ใช้'}"
                    )

                    last_id = fingerprint_id
                    last_time = now

            except Exception as exc:
                print(f"[SCANNER] ตรวจสอบลายนิ้วมือไม่สำเร็จ: {exc}")

            time.sleep(0.2)

    except Exception as exc:
        scanner_status["error"] = str(exc)
        print(f"[SCANNER ERROR] {exc}")

    finally:
        scanner_status["ready"] = False

        if opened:
            zk.CloseDevice()

        if initialized:
            zk.Terminate()

@app.route("/masters/subjects/import", methods=["POST"])
def subjects_import():
    file = request.files.get("subject_file")

    if not file or not file.filename:
        flash("กรุณาเลือกไฟล์ Excel", "error")
        return redirect(url_for("subjects_page"))

    if not file.filename.lower().endswith(".xlsx"):
        flash("รองรับเฉพาะไฟล์ Excel .xlsx", "error")
        return redirect(url_for("subjects_page"))

    try:
        workbook = load_workbook(
            BytesIO(file.read()),
            read_only=True,
            data_only=True
        )
        sheet = workbook.active

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            flash("ไฟล์ Excel ไม่มีข้อมูล", "error")
            return redirect(url_for("subjects_page"))

        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in headers
        ]

        required_columns = [
            "code",
            "name",
            "section_id",
            "class_days",
            "start_time",
            "end_time",
            "late_after_minutes",
        ]

        missing = [
            column for column in required_columns
            if column not in headers
        ]

        if missing:
            flash(
                "ไม่พบคอลัมน์ที่จำเป็น: " + ", ".join(missing),
                "error"
            )
            return redirect(url_for("subjects_page"))

        column_index = {
            name: headers.index(name)
            for name in required_columns
        }

        sections_result = (
            db.supabase.table("sections")
            .select("id")
            .execute()
        )
        valid_section_ids = {
            str(section["id"])
            for section in (sections_result.data or [])
        }

        existing_result = (
            db.supabase.table("subjects")
            .select("code")
            .execute()
        )
        existing_codes = {
            str(subject["code"]).strip()
            for subject in (existing_result.data or [])
            if subject.get("code")
        }

        def get_value(row, column):
            index = column_index[column]
            return row[index] if index < len(row) else None

        def format_time(value):
            if isinstance(value, datetime_time):
                return value.strftime("%H:%M:%S")

            if isinstance(value, datetime):
                return value.strftime("%H:%M:%S")

            if value is None:
                return None

            value = str(value).strip()

            if not value:
                return None

            for fmt in ("%H:%M", "%H:%M:%S"):
                try:
                    return datetime.strptime(value, fmt).strftime("%H:%M:%S")
                except ValueError:
                    continue

            raise ValueError(f"รูปแบบเวลาไม่ถูกต้อง: {value}")

        payloads = []
        errors = []
        seen_codes = set()

        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None for value in row):
                continue

            try:
                code = str(get_value(row, "code") or "").strip()
                name = str(get_value(row, "name") or "").strip()
                section_value = get_value(row, "section_id")
                class_days_value = get_value(row, "class_days")
                start_time_value = get_value(row, "start_time")
                end_time_value = get_value(row, "end_time")
                late_value = get_value(row, "late_after_minutes")

                if not code or not name:
                    raise ValueError("ต้องระบุ code และ name")

                if code in existing_codes or code in seen_codes:
                    raise ValueError(f"รหัสวิชาซ้ำ: {code}")

                if section_value is None:
                    raise ValueError("ต้องระบุ section_id")

                section_id = str(section_value).strip()

                if section_id.endswith(".0"):
                    section_id = section_id[:-2]

                if section_id not in valid_section_ids:
                    raise ValueError(
                        f"ไม่พบ Section ID: {section_id}"
                    )

                if not str(section_id).isdigit():
                    raise ValueError("section_id ต้องเป็นตัวเลข")

                if isinstance(class_days_value, (list, tuple)):
                    class_days = list(class_days_value)
                else:
                    class_days = [
                        day.strip()
                        for day in str(class_days_value or "").split(",")
                        if day.strip()
                    ]

                valid_days = {
                    "Monday", "Tuesday", "Wednesday", "Thursday",
                    "Friday", "Saturday", "Sunday"
                }

                if any(day not in valid_days for day in class_days):
                    raise ValueError(
                        "class_days ต้องใช้ชื่อวันภาษาอังกฤษ "
                        "คั่นด้วย comma"
                    )

                if late_value is None or str(late_value).strip() == "":
                    late_minutes = 0
                else:
                    late_minutes = int(late_value)

                if late_minutes < 0:
                    raise ValueError(
                        "late_after_minutes ต้องเป็น 0 ขึ้นไป"
                    )

                payloads.append({
                    "code": code,
                    "name": name,
                    "section_id": int(section_id),
                    "class_days": class_days,
                    "start_time": format_time(start_time_value),
                    "end_time": format_time(end_time_value),
                    "late_after_minutes": late_minutes,
                })

                seen_codes.add(code)

            except (ValueError, TypeError, IndexError) as exc:
                errors.append(f"แถว {row_number}: {exc}")

        workbook.close()

        if errors:
            preview = "\n".join(errors[:10])
            if len(errors) > 10:
                preview += f"\nและมีข้อผิดพลาดอีก {len(errors) - 10} แถว"

            flash(
                f"นำเข้าไม่สำเร็จ พบข้อผิดพลาด {len(errors)} แถว:\n{preview}",
                "error"
            )
            return redirect(url_for("subjects_page"))

        if not payloads:
            flash("ไม่พบข้อมูลรายวิชาสำหรับนำเข้า", "error")
            return redirect(url_for("subjects_page"))

        for start in range(0, len(payloads), 100):
            batch = payloads[start:start + 100]
            db.supabase.table("subjects").insert(batch).execute()

        flash(f"นำเข้ารายวิชาสำเร็จ {len(payloads)} รายการ", "success")

    except Exception as exc:
        print("Import subjects failed:", repr(exc))
        flash(f"นำเข้ารายวิชาไม่สำเร็จ: {exc}", "error")

    return redirect(url_for("subjects_page"))

scanner_thread = threading.Thread(
    target=scanner_worker,
    name="ZK9500Scanner",
    daemon=True
)
scanner_thread.start()

asgi_app = WsgiToAsgi(app)